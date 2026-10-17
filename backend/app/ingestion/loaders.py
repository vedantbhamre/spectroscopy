"""Page-aware document loaders: PDF · HTML · DOCX · PPTX · TXT · XLSX.

Each loader returns a list of (page, text). Pages are 1-based, matching how the
AFS materials are cited ("S01 p.10" = PDF page 10, counting the title page).
"""

import shutil
import subprocess
from pathlib import Path

Page = tuple[int, str]


def load_pdf(path: Path) -> list[Page]:
    """Prefer poppler's pdftotext (cleaner word spacing, so quotes match); fall back to pypdf."""
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    n = len(reader.pages)
    if shutil.which("pdftotext"):
        pages = []
        for i in range(1, n + 1):
            out = subprocess.run(["pdftotext", "-q", "-f", str(i), "-l", str(i), str(path), "-"], capture_output=True, text=True)
            pages.append((i, out.stdout))
        if any(t.strip() for _, t in pages):
            return pages
    return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]


def load_html(path: Path) -> list[Page]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    return [(1, soup.get_text("\n"))]


def load_docx(path: Path) -> list[Page]:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [(1, "\n".join(parts))]


def load_pptx(path: Path) -> list[Page]:
    from pptx import Presentation

    pages = []
    for i, slide in enumerate(Presentation(str(path)).slides):
        texts = [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        pages.append((i + 1, "\n".join(texts)))
    return pages


def load_txt(path: Path) -> list[Page]:
    return [(1, path.read_text(encoding="utf-8", errors="ignore"))]


def load_xlsx(path: Path) -> list[Page]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [" | ".join("" if c is None else str(c) for c in row) for row in ws.iter_rows(values_only=True)]
    return [(1, "\n".join(r for r in rows if r.strip(" |")))]


LOADERS = {
    ".pdf": load_pdf,
    ".html": load_html,
    ".htm": load_html,
    ".docx": load_docx,
    ".pptx": load_pptx,
    ".txt": load_txt,
    ".xlsx": load_xlsx,
}


def load(path: Path) -> list[Page]:
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise ValueError(f"Unsupported file type: {path.name}")
    return loader(path)
